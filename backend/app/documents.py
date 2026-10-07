"""上传文件的解析与批处理结果的还原。

- 文本类（TXT、Markdown、Word、PDF）：整份文件为一段，超过 CHUNK_CHARS 字时在换行处分段；
- 表格类（CSV 每行、JSON 每条）：一行（条）为一段，按“列名：值”逐行排列，所有列都参与识别，
  列名同时作为识别线索（“姓名”列的整格内容按姓名处理）；导出时还原成同样的列。
编码自动识别：UTF-8（含 BOM）、UTF-16（带 BOM）、GB18030（兼容 GBK，Excel 在中文 Windows 上的默认编码）。
"""
from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path
from typing import Any

from .schemas import EntityType, Span, Strategy

CHUNK_CHARS = 20_000
TEXT_SUFFIXES = {".txt", ".md", ".csv", ".json"}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | {".docx", ".pdf"}


class DocumentError(ValueError):
    """文件无法解析；message 直接展示给用户。status 为建议的 HTTP 状态码。"""

    def __init__(self, message: str, status: int = 400):
        super().__init__(message)
        self.status = status


def decode_text(filename: str, content: bytes) -> str:
    if content.startswith(b"\xef\xbb\xbf"):
        text = content[3:].decode("utf-8", errors="replace")
    elif content.startswith((b"\xff\xfe", b"\xfe\xff")):
        text = content.decode("utf-16", errors="replace")
    else:
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError:
            try:
                text = content.decode("gb18030")
            except UnicodeDecodeError as exc:
                raise DocumentError(f"{filename} 的文字编码无法识别，请用 UTF-8 编码另存后重试") from exc
    return text.replace("\r\n", "\n").replace("\r", "\n")


def chunk_text(text: str, limit: int = CHUNK_CHARS) -> list[str]:
    """在换行处把长文本切成不超过 limit 字的段；单行过长时在句末切开。拼接各段即得原文。"""
    if len(text) <= limit:
        return [text]
    pieces: list[str] = []
    for line in text.splitlines(keepends=True):
        while len(line) > limit:
            window = line[:limit]
            cut = max(window.rfind(mark) for mark in "。！？.!?；;")
            cut = cut + 1 if cut > limit // 2 else limit
            pieces.append(line[:cut])
            line = line[cut:]
        if line:
            pieces.append(line)
    chunks: list[str] = []
    for piece in pieces:
        if chunks and (len(chunks[-1]) + len(piece) <= limit or not piece.strip()):
            chunks[-1] += piece
        else:
            chunks.append(piece)
    return chunks


def _unique_keys(header: list[str]) -> list[str]:
    keys: list[str] = []
    for index, raw in enumerate(header):
        name = re.sub(r"\s+", " ", str(raw or "").replace("：", ":")).strip() or f"第{index + 1}列"
        candidate, suffix = name, 2
        while candidate in keys:
            candidate = f"{name}({suffix})"
            suffix += 1
        keys.append(candidate)
    return keys


def _table_record(filename: str, file_index: int, row: int, keys: list[str], values: list[str]) -> dict[str, Any] | None:
    lines: list[str] = []
    fields: list[dict[str, Any]] = []
    position = 0
    for key, value in zip(keys, values):
        if value is None or not str(value).strip():
            continue
        value = str(value).replace("\r\n", "\n").replace("\r", "\n")
        prefix = f"{key}："
        start = position + len(prefix)
        fields.append({"key": key, "start": start, "end": start + len(value)})
        lines.append(prefix + value)
        position += len(prefix) + len(value) + 1
    if not lines:
        return None
    return {"file": filename, "file_index": file_index, "row": row, "kind": "table", "text": "\n".join(lines), "fields": fields}


def _scalar_type(value: Any) -> str:
    if isinstance(value, bool):
        return "bool"
    if isinstance(value, int):
        return "int"
    if isinstance(value, float):
        return "float"
    if isinstance(value, str):
        return "str"
    return "json"


def parse_document(filename: str, content: bytes, file_index: int = 0) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """返回 (记录, 文件信息)。记录的 text 即一段待处理文本；文件信息用于导出时还原格式。"""
    suffix = Path(filename).suffix.lower()
    meta: dict[str, Any] = {"name": filename, "index": file_index, "format": suffix.lstrip(".") or "txt", "kind": "text"}
    if suffix not in SUPPORTED_SUFFIXES:
        raise DocumentError(f"暂不支持 {suffix or '无扩展名'} 文件：{filename}", 415)
    if suffix in TEXT_SUFFIXES:
        text = decode_text(filename, content)
        if suffix == ".csv":
            return _parse_csv(filename, text, file_index, meta)
        if suffix == ".json":
            return _parse_json(filename, text, file_index, meta)
        return _text_records(filename, text, file_index), meta
    if suffix == ".docx":
        return _text_records(filename, _docx_text(filename, content), file_index), meta
    return _text_records(filename, _pdf_text(filename, content), file_index), meta


def _text_records(filename: str, text: str, file_index: int) -> list[dict[str, Any]]:
    if not text.strip():
        return []
    return [{"file": filename, "file_index": file_index, "row": index + 1, "kind": "text", "text": chunk}
            for index, chunk in enumerate(chunk_text(text)) if chunk.strip()]


def _parse_csv(filename: str, text: str, file_index: int, meta: dict[str, Any]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    try:
        rows = [row for row in csv.reader(io.StringIO(text)) if any(cell.strip() for cell in row)]
    except csv.Error as exc:
        raise DocumentError(f"{filename} 不是有效的 CSV：{exc}") from exc
    if len(rows) < 2:
        raise DocumentError(f"{filename} 没有数据行（第一行应为列名）")
    keys = _unique_keys(rows[0])
    records = []
    for index, row in enumerate(rows[1:], start=1):
        cells = list(row) + [""] * (len(keys) - len(row))
        extra = cells[len(keys):]
        if any(cell.strip() for cell in extra):
            # 比列名多出来的格子并到最后一列，避免丢内容
            cells = cells[:len(keys) - 1] + [",".join(cells[len(keys) - 1:])]
        record = _table_record(filename, file_index, index, keys, cells[:len(keys)])
        if record:
            records.append(record)
    meta.update(kind="table", columns=keys)
    return records, meta


def _parse_json(filename: str, text: str, file_index: int, meta: dict[str, Any]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError as exc:
        raise DocumentError(f"JSON 格式错误（{filename} 第 {exc.lineno} 行）") from exc
    except RecursionError as exc:
        raise DocumentError(f"{filename} 的 JSON 层级过深，无法解析") from exc
    items = parsed if isinstance(parsed, list) else [parsed]
    records: list[dict[str, Any]] = []
    field_types: dict[str, str] = {}
    for index, item in enumerate(items, start=1):
        if item is None:
            continue
        if isinstance(item, dict):
            keys = _unique_keys([str(key) for key in item.keys()])
            values = []
            for key, value in zip(keys, item.values()):
                field_types.setdefault(key, _scalar_type(value))
                if value is None:
                    values.append("")
                elif isinstance(value, str):
                    values.append(value)
                elif isinstance(value, (dict, list)):
                    try:
                        values.append(json.dumps(value, ensure_ascii=False))
                    except RecursionError as exc:
                        raise DocumentError(f"{filename} 的 JSON 层级过深，无法解析") from exc
                else:
                    values.append(json.dumps(value, ensure_ascii=False))
            record = _table_record(filename, file_index, index, keys, values)
            if record:
                # 空值不含隐私，原样记下，导出时放回原位置
                record["keys"] = keys
                record["empty_values"] = {key: value for key, value in zip(keys, item.values()) if value is None or value == ""}
                records.append(record)
        else:
            value = item if isinstance(item, str) else json.dumps(item, ensure_ascii=False)
            if value.strip():
                records.append({"file": filename, "file_index": file_index, "row": index, "kind": "text", "text": value})
    meta.update(kind="table" if any(record["kind"] == "table" for record in records) else "text", field_types=field_types)
    return records, meta


def _docx_text(filename: str, content: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise DocumentError("服务器尚未安装 python-docx，无法读取 DOCX", 503) from exc
    try:
        document = Document(io.BytesIO(content))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            parts.append("")
            parts.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
    except Exception as exc:
        raise DocumentError(f"无法读取 DOCX 文件 {filename}，请确认文件完整且格式正确") from exc
    return "\n".join(parts).strip("\n").replace("\r\n", "\n").replace("\v", "\n")


def _pdf_text(filename: str, content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise DocumentError("服务器尚未安装 pypdf，无法读取 PDF", 503) from exc
    try:
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ValueError("PDF requires a password")
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
    except Exception as exc:
        raise DocumentError(f"无法读取 PDF 文件 {filename}，请确认文件完整、格式正确且未设置打开密码") from exc
    return "\n\n".join(page for page in pages if page).replace("\r\n", "\n")


def display_text(records: list[dict[str, Any]]) -> str:
    """导入工作台时的整段文字：文本类按原样拼回，表格类每行（条）之间空一行。"""
    parts: list[str] = []
    current: list[str] = []
    current_key: tuple[int, str] | None = None
    for record in records:
        key = (int(record.get("file_index", 0)), str(record.get("kind")))
        if key != current_key and current:
            parts.append(("" if current_key and current_key[1] == "text" else "\n\n").join(current))
            current = []
        current_key = key
        current.append(record["text"])
    if current:
        parts.append(("" if current_key and current_key[1] == "text" else "\n\n").join(current))
    return "\n\n".join(part.strip("\n") for part in parts)


# 列名提示：列名说明了这一格是什么，整格内容按对应类型处理（顺序重要：“联系人电话”先匹配电话）
_COLUMN_TYPES: list[tuple[EntityType, tuple[str, ...]]] = [
    (EntityType.EMAIL, ("邮箱", "电子邮件", "email", "e-mail")),
    (EntityType.PHONE, ("电话", "手机", "联系方式", "phone", "mobile", "tel")),
    (EntityType.ID_CARD, ("身份证", "证件号", "idcard", "id card", "id number")),
    (EntityType.BANK_CARD, ("银行卡", "卡号", "银行账号", "bank card", "card number")),
    (EntityType.PASSPORT, ("护照", "passport")),
    (EntityType.ADDRESS, ("地址", "住址", "address")),
    (EntityType.ORG, ("单位", "公司", "学校", "医院", "机构", "organization", "organisation", "company", "employer", "school", "hospital", "institution")),
    (EntityType.LOCATION, ("城市", "籍贯", "出生地", "所在地", "city", "hometown")),
    (EntityType.PERSON, ("姓名", "名字", "联系人", "客户", "患者", "学生", "员工", "负责人", "申请人", "收件人", "经办人", "name", "contact", "customer", "patient", "student", "employee", "applicant")),
]


def column_type(key: str) -> EntityType | None:
    lowered = key.lower()
    for entity_type, words in _COLUMN_TYPES:
        if any(word in lowered for word in words):
            return entity_type
    return None


def _plausible(entity_type: EntityType, value: str) -> bool:
    digits = sum(char.isdigit() for char in value)
    if entity_type == EntityType.PHONE:
        return digits >= 7
    if entity_type == EntityType.EMAIL:
        return "@" in value
    if entity_type == EntityType.ID_CARD:
        return 15 <= len(re.sub(r"\s", "", value)) <= 20 and digits >= 14
    if entity_type == EntityType.BANK_CARD:
        return digits >= 12
    if entity_type == EntityType.PASSPORT:
        return 5 <= len(value) <= 15
    if entity_type == EntityType.PERSON:
        return len(value) <= 30 and digits == 0
    return len(value) <= 200 and digits < len(value)


def column_hint_spans(record: dict[str, Any], strategy: Strategy) -> list[Span]:
    """表格记录中，按列名可以确定类型的整格内容。"""
    hints: list[Span] = []
    text = record.get("text", "")
    for field in record.get("fields") or []:
        entity_type = column_type(str(field.get("key", "")))
        start, end = int(field["start"]), int(field["end"])
        value = text[start:end]
        stripped = value.strip()
        if entity_type is None or not stripped or not _plausible(entity_type, stripped):
            continue
        start += len(value) - len(value.lstrip())
        end -= len(value) - len(value.rstrip())
        hints.append(Span(
            id=f"col_{record.get('file_index', 0)}_{record.get('row', 0)}_{start}", start=start, end=end, text=text[start:end],
            entity_type=entity_type, score=0.95, sources=["RULE"], status="accepted", strategy=strategy,
            metadata={"column": field.get("key"), "column_hint": True},
        ))
    return hints


def parse_table_text(text: str, keys: list[str]) -> dict[str, str]:
    """把“列名：值”逐行的文字还原成各列的值；值内的换行按续行处理。"""
    values: dict[str, str] = {}
    current: str | None = None
    known = sorted(keys, key=len, reverse=True)
    for line in text.split("\n"):
        key = next((item for item in known if line.startswith(f"{item}：")), None)
        if key is not None and key not in values:
            current = key
            values[key] = line[len(key) + 1:]
        elif current is not None:
            values[current] += "\n" + line
    return values


def restore_table(meta: dict[str, Any], rows: list[tuple[str, str]], row_meta: list[dict[str, Any]] | None = None) -> tuple[str, str]:
    """rows: [(原文, 最终稿)]；row_meta：每条的列顺序和空值（JSON）。返回 (导出文件后缀, 文件内容)。"""
    if meta.get("format") == "csv":
        columns = list(meta.get("columns") or [])
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(columns)
        for _, final in rows:
            values = parse_table_text(final, columns)
            writer.writerow([values.get(column, "") for column in columns])
        return ".csv", "﻿" + output.getvalue()
    field_types = meta.get("field_types") or {}
    items: list[Any] = []
    for position, (original, final) in enumerate(rows):
        extra = (row_meta or [{}] * len(rows))[position] or {}
        keys = list(extra.get("keys") or [])
        if not keys:
            keys = [line.split("：", 1)[0] for line in original.split("\n") if "：" in line]
            keys = [key for key in dict.fromkeys(keys) if key in field_types] or list(field_types)
        empty_values = extra.get("empty_values") or {}
        before = parse_table_text(original, keys)
        after = parse_table_text(final, keys)
        item: dict[str, Any] = {}
        for key in keys:
            if key in empty_values:
                item[key] = empty_values[key]
                continue
            if key not in after:
                continue
            value = after[key]
            kind = field_types.get(key, "str")
            if kind != "str" and value == before.get(key):
                try:
                    item[key] = json.loads(value)
                    continue
                except (ValueError, RecursionError):
                    pass
            item[key] = value
        items.append(item)
    return ".json", json.dumps(items, ensure_ascii=False, indent=2)
